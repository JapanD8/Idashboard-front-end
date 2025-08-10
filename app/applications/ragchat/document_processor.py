import os
import uuid
from typing import List, Dict, Any
from pathlib import Path
import PyPDF2
from docx import Document
from langchain_core.documents import Document as LangchainDocument
from langchain_google_genai import GoogleGenerativeAI, GoogleGenerativeAIEmbeddings
from langchain_core.prompts import PromptTemplate
from langchain_experimental.text_splitter import SemanticChunker
import re
import chromadb
from chromadb.config import Settings
import json
import google.generativeai as genai
import logging
from config import Config as config

GOOGLE_API_KEY = config.GOOGLE_API_KEY
SEMANTIC_MIN_CHUNK_SIZE = config.SEMANTIC_MIN_CHUNK_SIZE
SEMANTIC_MAX_CHUNK_SIZE = config.SEMANTIC_MAX_CHUNK_SIZE
SEMANTIC_BUFFER_SIZE = config.SEMANTIC_BUFFER_SIZE
SEMANTIC_THRESHOLD_TYPE = config.SEMANTIC_THRESHOLD_TYPE
SEMANTIC_THRESHOLD_AMOUNT = config.SEMANTIC_THRESHOLD_AMOUNT
SEMANTIC_NUMBER_OF_CHUNKS = config.SEMANTIC_NUMBER_OF_CHUNKS
# Set up logging
logger = logging.getLogger(__name__)

def clean_html_response(text: str) -> str:
    """Clean HTML tags while preserving text formatting."""
    if not text:
        return text
    
    # Remove HTML tags only
    clean_text = re.sub(r'<[^>]+>', '', text)
    
    # Clean up excessive blank lines but preserve formatting
    clean_text = re.sub(r'\n\s*\n\s*\n+', '\n\n', clean_text)
    
    # Remove leading/trailing whitespace only
    clean_text = clean_text.strip()
    
    return clean_text

# Configuration
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
UPLOAD_DIR = "uploads"

# # Ensure upload directory exists
# os.makedirs(UPLOAD_DIR, exist_ok=True)

# Initialize ChromaDB
chroma_client = chromadb.Client(Settings(anonymized_telemetry=False))
chroma_client = chromadb.PersistentClient(path="./chroma_db")
collection = chroma_client.get_or_create_collection(
    name="document_chunks",
    metadata={"hnsw:space": "cosine"}
)

# Configure Google AI
genai.configure(api_key=GOOGLE_API_KEY)

def create_text_chunks(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    """Simple text chunking function with overlap."""
    if len(text) <= chunk_size:
        return [text]
    
    chunks = []
    start = 0
    
    while start < len(text):
        end = start + chunk_size
        
        # Try to find a good breaking point
        if end < len(text):
            for sep in ["\n\n", "\n", ".", "!", "?", ";", ","]:
                break_point = text.rfind(sep, start, end)
                if break_point > start:
                    end = break_point + len(sep)
                    break
        
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        
        start = max(start + 1, end - overlap)
    
    return chunks

# Initialize embeddings (using Google's text-embedding model)
logger.debug("DEBUG: Initializing embeddings...")
try:
    embeddings = GoogleGenerativeAIEmbeddings(
        model="models/embedding-001",
        google_api_key=GOOGLE_API_KEY
    )
    logger.debug("DEBUG: Embeddings initialized successfully")
except Exception as e:
    logger.error(f"ERROR: Failed to initialize embeddings: {e}")
    raise


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from a PDF file."""
    text = ""
    try:
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
    except Exception as e:
        raise Exception(f"Error extracting text from PDF: {str(e)}")
    return text


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from a DOCX file."""
    try:
        doc = Document(file_path)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
    except Exception as e:
        raise Exception(f"Error extracting text from DOCX: {str(e)}")
    return text


def extract_text_from_doc(file_path: str) -> str:
    """Extract text from a DOC file (simplified - requires python-docx2txt or similar)."""
    # For now, we'll treat DOC files as text files
    # In production, you might want to use python-docx2txt or similar
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
            return file.read()
    except Exception as e:
        raise Exception(f"Error extracting text from DOC: {str(e)}")


def extract_text_from_file(file_path: str, file_type: str) -> str:
    """Extract text from a file based on its type."""
    file_type = file_type.lower()
    
    if file_type == "pdf":
        return extract_text_from_pdf(file_path)
    elif file_type == "docx":
        return extract_text_from_docx(file_path)
    elif file_type == "doc":
        return extract_text_from_doc(file_path)
    elif file_type == "txt":
        with open(file_path, 'r', encoding='utf-8') as file:
            return file.read()
    else:
        raise Exception(f"Unsupported file type: {file_type}")


def _split_large_chunks(chunks: List[str], max_size: int) -> List[str]:
    """Split chunks that exceed max_size into smaller chunks while preserving sentence boundaries."""
    result_chunks = []
    
    for chunk in chunks:
        if len(chunk) <= max_size:
            result_chunks.append(chunk)
        else:
            # Split large chunk by sentences first
            sentences = re.split(r'(?<=[.!?])\s+', chunk)
            current_chunk = ""
            
            for sentence in sentences:
                # If adding this sentence would exceed max_size, save current chunk and start new one
                if len(current_chunk) + len(sentence) + 1 > max_size:
                    if current_chunk:
                        result_chunks.append(current_chunk.strip())
                    current_chunk = sentence
                else:
                    current_chunk += (" " + sentence if current_chunk else sentence)
            
            # Add the remaining chunk
            if current_chunk:
                result_chunks.append(current_chunk.strip())
    
    return result_chunks

def create_semantic_chunks(text: str) -> List[str]:
    """Create semantic chunks from text using LangChain's SemanticChunker."""
    try:
        logger.debug(f"DEBUG create_semantic_chunks: Starting semantic chunking for text length: {len(text)}")
        
        # Initialize the SemanticChunker with Google embeddings
        semantic_splitter = SemanticChunker(
            embeddings,
            breakpoint_threshold_type=SEMANTIC_THRESHOLD_TYPE,  # percentile, standard_deviation, interquartile, gradient
            buffer_size=SEMANTIC_BUFFER_SIZE,  # Number of sentences to group together
            min_chunk_size=SEMANTIC_MIN_CHUNK_SIZE,  # Minimum chunk size in characters
            breakpoint_threshold_amount=SEMANTIC_THRESHOLD_AMOUNT,  # Custom threshold for breakpoints
            number_of_chunks=SEMANTIC_NUMBER_OF_CHUNKS,  # Target number of chunks (acts as max control)
        )
        
        # Create a LangChain document from the text
        langchain_doc = LangchainDocument(page_content=text)
        
        # Split the document semantically
        split_docs = semantic_splitter.split_documents([langchain_doc])
        
        # Extract the text content from the split documents
        chunks = [doc.page_content for doc in split_docs]
        
        # Post-process chunks to enforce max_size
        chunks = _split_large_chunks(chunks, SEMANTIC_MAX_CHUNK_SIZE)
        
        logger.debug(f"DEBUG create_semantic_chunks: Created {len(chunks)} semantic chunks")
        for i, chunk in enumerate(chunks[:3]):  # Log first 3 chunks as preview
            logger.debug(f"DEBUG create_semantic_chunks: Chunk {i+1} preview: {chunk[:100]}...")
        
        # If semantic chunking fails or produces no chunks, fall back to simple chunking
        if not chunks:
            logger.warning("DEBUG create_semantic_chunks: No chunks produced, falling back to simple chunking")
            return create_text_chunks(text)
        
        return chunks
        
    except Exception as e:
        logger.error(f"ERROR in create_semantic_chunks: {str(e)}")
        logger.error(f"ERROR create_semantic_chunks: Falling back to simple chunking")
        # Fall back to the simple text chunking if semantic chunking fails
        return create_text_chunks(text)


def generate_summary(text: str) -> str:
    """Generate a summary of the document text."""
    try:
        # Use Gemini to generate summary
        llm = GoogleGenerativeAI(
            model="gemini-2.0-flash-exp",
            google_api_key=GOOGLE_API_KEY,
            temperature=0
        )
        
        # Create a prompt for summarization
        prompt = f"""Please provide a concise summary of the following document in 2-3 sentences:

{text[:2000]}  # Limit text to avoid token limits

Summary:"""
        
        summary = llm.invoke(prompt)
        return clean_html_response(summary.strip())
    except Exception as e:
        # Fallback to first 200 characters if summarization fails
        print(f"Error generating summary: {str(e)}")
        return text[:200] + "..." if len(text) > 200 else text


def process_and_vectorize_document(
    file_path: str,
    file_id: str,
    user_id: int,
    file_name: str,
    file_type: str
) -> Dict[str, Any]:
    """Process a document and store it in the vector database."""
    try:
        # Extract text from file
        text = extract_text_from_file(file_path, file_type)
        
        if not text.strip():
            raise Exception("No text could be extracted from the file")
        
        # Generate summary
        summary = generate_summary(text)
        
        # Create semantic chunks
        chunks = create_semantic_chunks(text)
        
        # Generate embeddings and store in ChromaDB
        chunk_ids = []
        chunk_data = []
        
        for i, chunk in enumerate(chunks):
            chunk_id = f"{file_id}_chunk_{i}"
            chunk_ids.append(chunk_id)
            
            # Create embedding
            embedding = embeddings.embed_query(chunk)
            
            # Prepare metadata
            metadata = {
                "file_id": file_id,
                "user_id": str(user_id),
                "file_name": file_name,
                "file_type": file_type,
                "chunk_index": i,
                "summary": summary
            }
            
            chunk_data.append({
                "id": chunk_id,
                "embedding": embedding,
                "document": chunk,
                "metadata": metadata
            })
        
        # Add to ChromaDB collection
        collection.add(
            ids=[item["id"] for item in chunk_data],
            embeddings=[item["embedding"] for item in chunk_data],
            documents=[item["document"] for item in chunk_data],
            metadatas=[item["metadata"] for item in chunk_data]
        )
        
        return {
            "success": True,
            "summary": summary,
            "chunk_count": len(chunks),
            "chunk_ids": chunk_ids
        }
        
    except Exception as e:
        return {
            "success": False,
            "error": str(e)
        }


def search_documents(query: str, user_id: int, file_ids: List[str] = None, top_k: int = 10) -> List[Dict]:
    """Search for relevant document chunks."""
    try:
        logger.debug(f"DEBUG search_documents: Starting search")
        logger.debug(f"DEBUG search_documents: Query: {query}")
        logger.debug(f"DEBUG search_documents: User ID: {user_id}")
        logger.debug(f"DEBUG search_documents: File IDs: {file_ids}")
        logger.debug(f"DEBUG search_documents: Top K: {top_k}")
        
        # Generate query embedding
        logger.debug(f"DEBUG search_documents: Generating query embedding...")
        query_embedding = embeddings.embed_query(query)
        logger.debug(f"DEBUG search_documents: Query embedding length: {len(query_embedding)}")
        
        # Prepare filter
        if file_ids:
            where_filter = {
                "$and": [
                    {"user_id": str(user_id)},
                    {"file_id": {"$in": file_ids}}
                ]
            }
        else:
            where_filter = {"user_id": str(user_id)}
        
        logger.debug(f"DEBUG search_documents: Where filter: {where_filter}")
        
        # First, let's check what's actually in the collection
        logger.debug(f"DEBUG search_documents: Checking collection contents...")
        try:
            # Get a sample of all documents in the collection
            all_results = collection.get(
                where={"user_id": str(user_id)},
                limit=100
            )
            logger.debug(f"DEBUG search_documents: Found {len(all_results['ids'])} total chunks for user")
            
            if all_results['ids']:
                for i, doc_id in enumerate(all_results['ids'][:3]):  # Show first 3
                    metadata = all_results['metadatas'][i] if all_results['metadatas'] else {}
                    logger.debug(f"DEBUG search_documents: Sample chunk {i+1}: ID={doc_id}, metadata={metadata}")
            
            # If file_ids specified, check if those files exist
            if file_ids:
                for file_id in file_ids:
                    file_chunks = collection.get(
                        where={"file_id": file_id}
                    )
                    logger.debug(f"DEBUG search_documents: File {file_id} has {len(file_chunks['ids'])} chunks")
                    
        except Exception as e:
            logger.error(f"DEBUG search_documents: Error checking collection: {e}")
        
        # Search in ChromaDB
        logger.debug(f"DEBUG search_documents: Performing vector search...")
        results = collection.query(
            query_embeddings=[query_embedding],
            n_results=top_k,
            where=where_filter
        )
        
        logger.debug(f"DEBUG search_documents: Raw results structure:")
        logger.debug(f"  - IDs: {results.get('ids', [])}")
        logger.debug(f"  - Documents: {len(results.get('documents', [[]]))} docs")
        logger.debug(f"  - Metadatas: {len(results.get('metadatas', [[]]))} metadata entries")
        logger.debug(f"  - Distances: {results.get('distances', [])}")
        
        # Format results
        formatted_results = []
        if results["documents"] and len(results["documents"]) > 0 and len(results["documents"][0]) > 0:
            logger.debug(f"DEBUG search_documents: Processing {len(results['documents'][0])} search results")
            for i in range(len(results["documents"][0])):
                result_data = {
                    "chunk_id": results["ids"][0][i],
                    "content": results["documents"][0][i],
                    "metadata": results["metadatas"][0][i],
                    "similarity": 1 - results["distances"][0][i]  # Convert distance to similarity
                }
                formatted_results.append(result_data)
                logger.debug(f"DEBUG search_documents: Formatted result {i+1}:")
                logger.debug(f"  - Chunk ID: {result_data['chunk_id']}")
                logger.debug(f"  - Similarity: {result_data['similarity']:.4f}")
                logger.debug(f"  - Content preview: {result_data['content'][:100]}...")
        else:
            logger.warning(f"DEBUG search_documents: No documents found in results!")
        
        logger.debug(f"DEBUG search_documents: Returning {len(formatted_results)} formatted results")
        return formatted_results
        
    except Exception as e:
        logger.error(f"ERROR in search_documents: {str(e)}")
        logger.error(f"ERROR search_documents details: {repr(e)}")
        import traceback
        logger.error(f"ERROR search_documents traceback: {traceback.format_exc()}")
        print(f"Error searching documents: {str(e)}")
        return []


def delete_document_chunks(file_id: str):
    """Delete all chunks for a specific file from the vector database."""
    try:
        # Get all chunk IDs for this file
        results = collection.get(where={"file_id": file_id})
        if results["ids"]:
            collection.delete(ids=results["ids"])
        return True
    except Exception as e:
        print(f"Error deleting document chunks: {str(e)}")
        return False 